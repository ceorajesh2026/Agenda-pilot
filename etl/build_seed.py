# -*- coding: utf-8 -*-
"""AgendaPilot ETL — parse the real WNS 2026 source files into a normalized seed.json.

Pipeline:  program xlsx + 2 neurosurgery docx + faculty directory
   -> Event/Track/Day/Hall/Session/Slot/Role/Person  (+ entity resolution)
Output: app/src/data/seed.json  and  etl/build/resolution-review.csv
"""
import openpyxl, docx, re, json, os, difflib, csv
from directory import load_directory, norm_name, slugify

BASE = r'C:/ForgeAI_residency/Agenda_management/'
APP_DATA = os.path.join(BASE, 'app', 'src', 'data')
BUILD = os.path.join(BASE, 'etl', 'build')
os.makedirs(APP_DATA, exist_ok=True)
os.makedirs(BUILD, exist_ok=True)

WEEKDAY_DATE = {
    'THURSDAY': '2026-08-20', 'FRIDAY': '2026-08-21',
    'SATURDAY': '2026-08-22', 'SUNDAY': '2026-08-23',
}
DAY_LABEL = {
    '2026-08-20': 'Pre-Conference (Thu 20 Aug)', '2026-08-21': 'Day 1 (Fri 21 Aug)',
    '2026-08-22': 'Day 2 (Sat 22 Aug)', '2026-08-23': 'Day 3 (Sun 23 Aug)',
}

ROLE_LABELS = [
    ('grand master', 'grand_master'), ('quiz master', 'quiz_master'),
    ('session expert', 'session_expert'), ('workshop director', 'workshop_director'),
    ('clinical expert', 'clinical_expert'), ('pathology expert', 'pathology_expert'),
    ('expert panellist', 'panellist'), ('expert panelist', 'panellist'),
    ('panellist', 'panellist'), ('panelist', 'panellist'),
    ('moderator', 'moderator'), ('expert', 'session_expert'),
]
SECTION_WORDS = {
    'BREAKFAST SESSION', 'PRE LUNCH SESSIONS', 'POST LUNCH SESSIONS', 'POST LUNCH SESSION',
    'SUNSET SESSION', 'SESSION 1', 'SESSION 2', 'SESSION 3',
    'PARALLEL NEUROLOGY WORKSHOPS / SYMPOSIUM', 'SCIENTIFIC PROGRAM - NEUROLOGY',
}
BREAK_WORDS = ('TEA BREAK', 'LUNCH', 'COCKTAILS', 'DINNER', 'BREAK')
LOCK_HINTS = ('KEYNOTE', 'INAUGURAL', 'MEMORIAL QUIZ', 'ASHOK PANAGARIYA')

TIME_RANGE = re.compile(r'(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})')
DUR = re.compile(r'\((\d{1,2})\s*min', re.I)


def to_min(h, m):
    return int(h) * 60 + int(m)


def parse_time_range(cell):
    if not cell:
        return None, None
    mt = TIME_RANGE.search(str(cell).replace(' ', ' '))
    if not mt:
        return None, None
    return to_min(mt.group(1), mt.group(2)), to_min(mt.group(3), mt.group(4))


def hhmm(mins):
    return None if mins is None else f'{mins // 60:02d}:{mins % 60:02d}'


def strip_dur(title):
    return re.sub(r'\s*[/(]?\s*\(\d{1,2}\s*min[^)]*\)\s*', ' ', title, flags=re.I).strip(' /')


def split_people(cell):
    """Split a speaker cell into individual names (handles '/' and trailing commas)."""
    if not cell:
        return []
    parts = re.split(r'[/\n]', str(cell))
    out = []
    for p in parts:
        p = p.strip(' ,')
        if p:
            out.append(p)
    return out


def parse_role_line(text):
    """'Moderator: A, B / Session Experts: C, D, / E' -> [(role_type, [names...])]"""
    frags = re.split(r'[/\n]', text)
    groups = []
    cur = None
    for frag in frags:
        frag = frag.strip()
        if not frag:
            continue
        low = frag.lower()
        matched = None
        for label, rtype in ROLE_LABELS:
            idx = low.find(label)
            if idx != -1 and (':' in frag or idx == 0):
                colon = frag.find(':')
                names_part = frag[colon + 1:] if colon != -1 else frag[idx + len(label):]
                matched = (rtype, names_part)
                break
        if matched:
            cur = {'role': matched[0], 'names': []}
            groups.append(cur)
            names_part = matched[1]
        else:
            names_part = frag
        if cur is not None:
            for nm in re.split(r',', names_part):
                nm = nm.strip(' ,')
                if nm and not any(lbl in nm.lower() for lbl, _ in ROLE_LABELS):
                    cur['names'].append(nm)
    return [(g['role'], g['names']) for g in groups if g['names']]


def is_role_line(text):
    low = text.lower()
    return ':' in text and any(lbl in low for lbl, _ in ROLE_LABELS)


# ------------------------------------------------------------------ program xlsx
def parse_program():
    wb = openpyxl.load_workbook(BASE + 'Conference Program Schedule.xlsx', data_only=True, read_only=True)
    ws = wb['Sheet1']
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    sessions, slots, roles = [], [], []
    cur_date = None
    parallel = False
    cur_session = None
    umbrella = None   # long container session (workshop/symposium) that spans breaks
    sid = tid = rid = 0

    def cell(r, i):
        return str(r[i]).strip() if (len(r) > i and r[i] is not None and str(r[i]).strip()) else ''

    for r in rows:
        c1, c2, c3 = cell(r, 1), cell(r, 2), cell(r, 3)
        joined = f'{c1} {c2}'.upper()

        # day header (weekday name present)
        wd = next((w for w in WEEKDAY_DATE if w in joined), None)
        if wd and ('AUGUST' in joined or 'DAY' in joined):
            cur_date = WEEKDAY_DATE[wd]
            cur_session = umbrella = None
            continue
        if 'PARALLEL NEUROLOGY WORKSHOPS' in joined:
            parallel = True
            cur_session = umbrella = None
            continue

        start, end = parse_time_range(c1)

        # session-start row: has a time range
        if start is not None and c2:
            title = strip_dur(c2)
            up = c2.upper()
            stype = 'break' if any(b in up for b in BREAK_WORDS) else 'session'
            locked = any(h in up for h in LOCK_HINTS)

            # A timed row that falls inside an active umbrella (a long workshop/symposium
            # container) is a SLOT of that umbrella, not a new parallel session. The
            # umbrella persists across TEA BREAK / LUNCH rows.
            inside = (umbrella is not None and umbrella['end_min'] is not None
                      and umbrella['start_min'] <= start < umbrella['end_min'])
            if inside:
                cur_session = umbrella
                if stype == 'break':
                    continue  # a break inside a workshop: not a schedulable conflict
                mdur = DUR.search(c2)
                dur = int(mdur.group(1)) if mdur else (end - start if end else None)
                kind = 'qa' if re.match(r'\s*q\s*&?\s*a', c2, re.I) else 'talk'
                tid += 1
                slot = {'id': f'sl{tid}', 'session_id': umbrella['id'], 'title': title,
                        'duration_min': dur, 'kind': kind,
                        'order': len([s for s in slots if s['session_id'] == umbrella['id']])}
                slots.append(slot)
                for nm in split_people(c3):
                    rid += 1
                    roles.append({'id': f'r{rid}', 'session_id': umbrella['id'],
                                  'slot_id': slot['id'], 'role_type': 'speaker', 'name_raw': nm})
                continue

            sid += 1
            cur_session = {
                'id': f's{sid}', 'title': title, 'type': stype,
                'date': cur_date, 'start_min': start, 'end_min': end,
                'start': hhmm(start), 'end': hhmm(end),
                'track': 'neurology', 'stream': 'parallel-workshop' if parallel else 'main',
                'state': 'PUBLISHED', 'locked': locked,
            }
            sessions.append(cur_session)
            if stype == 'session' and not c3:
                umbrella = cur_session          # a container: following timed talks nest here
            else:
                umbrella = None
                if c3 and stype != 'break':     # atomic standalone talk (time + title + speaker)
                    tid += 1
                    mdur = DUR.search(c2)
                    slots.append({'id': f'sl{tid}', 'session_id': cur_session['id'],
                                  'title': title, 'duration_min': int(mdur.group(1)) if mdur else None,
                                  'kind': 'talk', 'order': 0})
                    for nm in split_people(c3):
                        rid += 1
                        roles.append({'id': f'r{rid}', 'session_id': cur_session['id'],
                                      'slot_id': f'sl{tid}', 'role_type': 'speaker', 'name_raw': nm})
            continue

        # sub-rows (no time) belong to current session
        if cur_session is None:
            continue
        content = c2
        if not content:
            continue
        up = content.upper()
        if up in SECTION_WORDS or (len(content) < 24 and up == content and ':' not in content and not c3):
            continue  # section grouping label

        if is_role_line(content):
            for rtype, names in parse_role_line(content):
                for nm in names:
                    rid += 1
                    roles.append({'id': f'r{rid}', 'session_id': cur_session['id'],
                                  'slot_id': None, 'role_type': rtype, 'name_raw': nm})
            continue

        # otherwise it's a talk / Q&A slot
        mdur = DUR.search(content)
        kind = 'qa' if content.lower().startswith('q') and 'a' in content.lower()[:6] else 'talk'
        tid += 1
        slot = {'id': f'sl{tid}', 'session_id': cur_session['id'], 'title': strip_dur(content),
                'duration_min': int(mdur.group(1)) if mdur else None, 'kind': kind,
                'order': len([s for s in slots if s['session_id'] == cur_session['id']])}
        slots.append(slot)
        for nm in split_people(c3):
            rid += 1
            roles.append({'id': f'r{rid}', 'session_id': cur_session['id'],
                          'slot_id': slot['id'], 'role_type': 'speaker', 'name_raw': nm})
    return sessions, slots, roles


# ------------------------------------------------------------------ neurosurgery docx
def parse_ns(path, track_name):
    d = docx.Document(path)
    sessions, slots = [], []
    cur_date = None
    cur = None
    sid = tid = 0
    prefix = re.sub(r'[^a-z]+', '', track_name.lower())[:4]
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        up = t.upper()
        wd = next((w for w in WEEKDAY_DATE if w in up), None)
        if wd and 'AUGUST' in up:
            cur_date = WEEKDAY_DATE[wd]
            cur = None
            continue
        # session markers
        if re.match(r'^Session\s+[A-Z]\s*:', t) or re.match(r'^(Breakfast|Morning|Afternoon)', t, re.I) \
                or re.match(r'^Session\s+[A-Z]\b', t):
            sid += 1
            band = None
            if re.search(r'morning', t, re.I):
                band = 'morning'
            elif re.search(r'afternoon', t, re.I):
                band = 'afternoon'
            elif re.search(r'breakfast', t, re.I):
                band = 'breakfast'
            cur = {'id': f'{prefix}s{sid}', 'title': t, 'type': 'session', 'date': cur_date,
                   'start_min': None, 'end_min': None, 'start': None, 'end': None,
                   'track': track_name, 'stream': 'main', 'state': 'DRAFT', 'locked': False,
                   'band': band}
            sessions.append(cur)
            continue
        if cur is None:  # topic before any session header -> make a generic session
            sid += 1
            cur = {'id': f'{prefix}s{sid}', 'title': 'Unsorted', 'type': 'session', 'date': cur_date,
                   'start_min': None, 'end_min': None, 'start': None, 'end': None,
                   'track': track_name, 'stream': 'main', 'state': 'DRAFT', 'locked': False, 'band': None}
            sessions.append(cur)
        tid += 1
        slots.append({'id': f'{prefix}sl{tid}', 'session_id': cur['id'], 'title': t,
                      'duration_min': None, 'kind': 'talk', 'order': 0})
    return sessions, slots


# ------------------------------------------------------------------ entity resolution
def resolve(roles, people):
    by_key = {}
    for p in people:
        by_key.setdefault(p['name_key'], p)
    nospace = {k.replace(' ', ''): p for k, p in by_key.items()}
    keys = list(by_key.keys())
    review = []
    stats = {'exact': 0, 'fuzzy': 0, 'ambiguous': 0, 'unmatched': 0, 'group': 0}
    for role in roles:
        raw = role['name_raw']
        if raw.strip().lower() in ('all faculty', 'all faculties'):
            role['person_id'] = None
            role['match'] = 'group'
            stats['group'] += 1
            continue
        key = norm_name(raw)
        p = by_key.get(key) or nospace.get(key.replace(' ', ''))
        if p:
            role['person_id'] = p['id']
            role['match'] = 'exact'
            stats['exact'] += 1
            continue
        cand = difflib.get_close_matches(key, keys, n=3, cutoff=0.86)
        if len(cand) == 1:
            p = by_key[cand[0]]
            role['person_id'] = p['id']
            role['match'] = 'fuzzy'
            stats['fuzzy'] += 1
            review.append((raw, key, 'FUZZY -> ' + p['name'], f"{difflib.SequenceMatcher(None, key, cand[0]).ratio():.2f}"))
        elif len(cand) > 1:
            role['person_id'] = None
            role['match'] = 'ambiguous'
            stats['ambiguous'] += 1
            review.append((raw, key, 'AMBIGUOUS: ' + ' | '.join(cand), ''))
        else:
            role['person_id'] = None
            role['match'] = 'unmatched'
            stats['unmatched'] += 1
            review.append((raw, key, 'UNMATCHED', ''))
    return review, stats


# ------------------------------------------------------------------ hall inference
def assign_halls(sessions):
    halls = {}
    def hall(name):
        hid = 'h-' + re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')
        halls.setdefault(hid, {'id': hid, 'name': name, 'provisional': True})
        return hid
    for s in sessions:
        if s['track'] == 'neurology' and s['stream'] == 'main':
            s['hall_id'] = hall('Hall A — Neurology Main')
        elif s['track'] == 'neurology':
            # Parallel workshops run concurrently in real life -> each its own provisional
            # room, so the engine does not fabricate hall clashes on unknown room data.
            s['hall_id'] = hall('Room: ' + (s['title'][:34]))
        elif 'SPINE' in s['track'].upper():
            s['hall_id'] = hall('Hall C — Spine Olympiad')
        else:
            s['hall_id'] = hall('Hall D — Neuro-Odyssey')
    return list(halls.values())


# ------------------------------------------------------------------ main
def main():
    people = load_directory(BASE)
    n_sessions, n_slots, n_roles = parse_program()
    sp_s, sp_sl = parse_ns(BASE + 'WNS 2026 - SPINE OLYMPIAD PROGRAM.docx', 'Spine Olympiad')
    no_s, no_sl = parse_ns(BASE + 'WNS 2026 - THE NEURO-ODYSSEY PROGRAM.docx', 'Neuro-Odyssey')

    sessions = n_sessions + sp_s + no_s
    slots = n_slots + sp_sl + no_sl
    roles = n_roles  # only neurology program has named roles

    review, stats = resolve(roles, people)
    halls = assign_halls(sessions)

    days = [{'date': d, 'label': DAY_LABEL[d]} for d in sorted({s['date'] for s in sessions if s['date']})]
    tracks = [
        {'id': 'neurology', 'name': 'Neurology', 'kind': 'neurology', 'status': 'finalized'},
        {'id': 'spine', 'name': 'Spine Olympiad', 'kind': 'neurosurgery', 'status': 'tentative'},
        {'id': 'odyssey', 'name': 'Neuro-Odyssey', 'kind': 'neurosurgery', 'status': 'tentative'},
    ]

    seed = {
        'event': {'id': 'wns2026', 'name': 'World Neurosciences Summit 2026',
                  'start_date': '2026-08-20', 'end_date': '2026-08-23',
                  'website': 'https://www.worldneurosciencessummit.com/'},
        'generated_note': 'Auto-generated by etl/build_seed.py from the source xlsx/docx. '
                          'Halls are provisional (inferred). NS tracks are DRAFT (no times/speakers in source).',
        'tracks': tracks, 'days': days, 'halls': halls,
        'people': sorted(people, key=lambda p: p['name']),
        'sessions': sessions, 'slots': slots, 'roles': roles,
        'resolution_review': [{'raw': a, 'key': b, 'result': c, 'score': d} for a, b, c, d in review],
        'resolution_stats': stats,
    }

    with open(os.path.join(APP_DATA, 'seed.json'), 'w', encoding='utf-8') as f:
        json.dump(seed, f, ensure_ascii=False, indent=1)
    with open(os.path.join(BUILD, 'resolution-review.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['raw_name', 'norm_key', 'result', 'score'])
        w.writerows(review)

    # ---- console summary ----
    nns = [s for s in sessions if s['track'] == 'neurology']
    print('=== SEED BUILD SUMMARY ===')
    print(f'People (deduped)      : {len(people)}')
    print(f'Tracks                : {len(tracks)}  | Days: {len(days)} | Halls(provisional): {len(halls)}')
    print(f'Sessions total        : {len(sessions)}  (neurology {len(nns)}, '
          f'spine {len(sp_s)}, odyssey {len(no_s)})')
    print(f'Slots total           : {len(slots)}')
    print(f'Roles parsed (neuro)  : {len(roles)}')
    print(f'Entity resolution     : {stats}')
    total = len(roles) or 1
    print(f'  resolved            : {stats["exact"]+stats["fuzzy"]}/{len(roles)} '
          f'({100*(stats["exact"]+stats["fuzzy"])//total}%)  '
          f'needs review: {stats["ambiguous"]+stats["unmatched"]}  group(All Faculty): {stats["group"]}')
    ns_slots = len(sp_sl) + len(no_sl)
    print(f'Unassigned NS slots   : {ns_slots} (DRAFT, no time/speaker)')
    print(f'WROTE: {os.path.join(APP_DATA, "seed.json")}')
    print(f'WROTE: {os.path.join(BUILD, "resolution-review.csv")}')


if __name__ == '__main__':
    main()
